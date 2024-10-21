import re
import sys
from typing import List

import json

import pandas as pd
import pymupdf

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows


def clean_annot_in_doc(doc):
    # remove annotation information from pdf files
    # to avoid the impact of annotation information on form extraction
    for page in doc:
        for annot in page.annots():
            page.delete_annot(annot=annot)


def sort_size_list(size_set):
    size_to_number = {
        "XXXS": 1,
        "XXS": 2,
        "XS": 3,
        "S": 4,
        "M": 5,
        "L": 6,
        "XL": 7,
        "XXL": 8,
        "XXXL": 9,
    }

    def sort_sizes_str(sizes):
        sorted_sizes = sorted(
            sizes, key=lambda size: size_to_number.get(size, 0), reverse=False
        )
        return sorted_sizes

    bool_all_number = True
    size_list = []
    for tmp_size_set in size_set:
        try:
            # tmp_size_set = int(tmp_size_set)
            size_list.append(int(tmp_size_set))
        except Exception as e:
            size_list.append(tmp_size_set)
            bool_all_number = False

    # for tmp_size in size_list:
    #     print(f"tmp_size:{tmp_size}")
    # print(f"bool_all_number:{bool_all_number}")
    if bool_all_number is False:
        sorted_size = sort_sizes_str(size_list)
    else:
        sorted_size = sorted(size_list, reverse=False)

    return sorted_size


def trans_json2ws(total_style_info_list, size_columns_set):
    new_df = pd.DataFrame(total_style_info_list)
    new_df.rename(
        columns={
            "ITEM": "款号",
            "COLOUR": "颜色",
            "COST USD": "单价",
            "QTY": "数量",
            "TOTAL USD": "总金额",
        },
        inplace=True,
    )

    new_df["单价"] = new_df["单价"].astype(float)
    new_df["数量"] = new_df["数量"].astype(float)
    new_df["总金额"] = new_df["总金额"].astype(float)

    front_list = ["PO号", "款号", "颜色"]
    end_list = ["数量", "单价", "总金额", "离厂时间"]

    size_list = sort_size_list(size_set=size_columns_set)

    print(f"size_list:{size_list}")
    new_order = front_list + size_list + end_list
    exist_index = new_df.columns
    print(f"new_df.columns:{new_df.columns}")
    for ordered_key in new_order:
        if ordered_key in exist_index:
            pass
        else:
            new_df[ordered_key] = pd.NA
    new_df = new_df[new_order]

    if "DESCRIPTION" in new_df.columns:
        new_df.drop(columns="DESCRIPTION", inplace=True)

    print(f"new_df:{new_df}")

    wb = Workbook()

    # 获取当前活跃的工作表
    ws = wb.active

    # 将DataFrame的数据写入工作表
    for r in dataframe_to_rows(new_df, index=False, header=True):
        print(f"r:{r}")
        ws.append(r)
    return wb


def func_pdf2excel(pdf_content, size_columns_set=set()):
    """transform pdf to excel

    Args:
        pdf_content (Union[str, bytes]): file path or bytes
        size_columns_set (set, optional): the set of size columns. Defaults to set().

    Returns:
        _type_: _description_
    """
    if isinstance(pdf_content, str):
        doc = pymupdf.open(pdf_content)
    else:
        doc = pymupdf.open(stream=pdf_content)

    clean_annot_in_doc(doc=doc)

    # find target block content
    tar_content_list = ["camilla and marc Order No:", "EX FACTORY DATE:"]
    search_content_res_list = []
    page = doc[0]
    page_content_list = get_page_content(page=page)
    print(f"page_content_list:{page_content_list}")
    for tmp_tar_content_list in tar_content_list:
        bool_find = False
        for tmp_page_content_list in page_content_list:
            if tmp_tar_content_list in tmp_page_content_list:
                search_content_res_list.append(tmp_page_content_list)
                bool_find = True
                break

        if bool_find is False:
            search_content_res_list.append("")

    # get target info in target block content
    target_content_res_list = []
    for tmp_search_content_res in search_content_res_list:
        tmp_target_content_res = tmp_search_content_res.split(":")[-1].strip()
        target_content_res_list.append(tmp_target_content_res)

    PO = target_content_res_list[0]
    time = target_content_res_list[1]

    tables = page.find_tables()
    print("get tables")
    for table in tables:
        table_df = table.to_pandas()
        table_json = json.loads(table_df.to_json(orient="records", force_ascii=False))
        print("get table")

        # get size cols
        cols_size = {}
        table_infos_json_list = []
        for row_idx, tmp_table_row in enumerate(table_json):
            if row_idx == 0:
                for key, value in tmp_table_row.items():
                    if value:
                        cols_size[key] = int(value)
            else:
                tmp_table_infos_json = {
                    "PO号": PO,
                    "离厂时间": time,
                }
                for key, value in tmp_table_row.items():
                    if value:
                        if key in cols_size:
                            tmp_table_infos_json[cols_size[key]] = int(value)
                        else:
                            tmp_table_infos_json[key] = value

                table_infos_json_list.append(tmp_table_infos_json)
        print(f"table_infos_json_list:{table_infos_json_list}")
        break

    # size_columns_set = set()
    for key, val in cols_size.items():
        size_columns_set.add(val)
    if isinstance(pdf_content, str):
        wb = trans_json2ws(
            total_style_info_list=table_infos_json_list,
            size_columns_set=size_columns_set,
        )
        wb.save("a.xlsx")
        # new_df.to_excel("a.xlsx")
    else:
        return table_infos_json_list, size_columns_set


def mark_pdf(input_path: str, output_path: str = "./", level: str = ""):
    """mark PDF content with rectangle

    Args:
        input_path (str): input file path
        output_path (str, optional): output file folder path. Defaults to "./".
        level (str, optional): mark content tags, include block, line, span, table, cell. Defaults to "".
    """
    doc = pymupdf.open(input_path)

    for page in doc:
        page_content = page.get_text(option="dict")

        for block in page_content["blocks"]:
            if level == "block":
                page.draw_rect(pymupdf.Rect(block["bbox"]), color=(1, 0, 0))

            if "lines" in block:
                for line in block["lines"]:
                    if level == "line":
                        page.draw_rect(pymupdf.Rect(line["bbox"]))
                    for span in line["spans"]:
                        if level == "span":
                            page.draw_rect(pymupdf.Rect(span["bbox"]))

        if level in ["table", "cell"]:
            tables = page.find_tables()
            print("get tables")

            if level == "cell":
                for table in tables.tables:
                    for cell in table.cells:
                        table.page.draw_rect(cell, color=(1, 0, 0))

            if level == "table":
                for table in tables.tables:
                    table.page.draw_rect(table.bbox, color=(0, 1, 0))

    doc.save(f"{output_path}/{level}.pdf")


def get_page_content(page) -> List[str]:
    """get page content in PDF

    Args:
        page (_type_): page object

    Returns:
        List[str]: content list
    """

    page_content = page.get_text(option="dict")
    total_block_content_list = []
    for block in page_content["blocks"]:
        tmp_block_content = ""
        tmp_block_list = []
        if "lines" in block:
            for line in block["lines"]:
                for span in line["spans"]:
                    tmp_block_list.append(span["text"])

        tmp_block_content = " ".join(tmp_block_list)
        total_block_content_list.append(tmp_block_content)

    return total_block_content_list


if __name__ == "__main__":
    ORG_PDF_PATH = "D:/projects/pdf2excel/pdf2excel_GAUDI_camilla/others/sample_file/P018430 - LATEEN MINI DRESS O2CMD 1738.DBLK.pdf"

    # mark_pdf(input_path=ORG_PDF_PATH, output_path="./P018430", level="cell")
    # total_style_info_list, size_columns_set = func_pdf2excel(pdf_content=ORG_PDF_PATH)
    pdf_content = ORG_PDF_PATH
    func_pdf2excel(pdf_content=pdf_content)
